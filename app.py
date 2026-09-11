from fastapi import FastAPI, Request, Form
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware
import psycopg2
from psycopg2.extras import RealDictCursor
import pandas as pd
from io import BytesIO
import os
import hashlib
import uuid
from datetime import datetime
from dotenv import load_dotenv
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

load_dotenv()

app = FastAPI(title="Programas Bienestar")
app.add_middleware(SessionMiddleware, secret_key="bienestar_super_secreto_2026")
templates = Jinja2Templates(directory="templates")

DATABASE_URL = os.getenv("DATABASE_URL")

def hash_password(password: str):
    return hashlib.sha256(password.encode()).hexdigest()

def get_db_connection():
    return psycopg2.connect(DATABASE_URL)

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()
    
    # 1. Tabla de beneficiarios
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS beneficiarios (
            id SERIAL PRIMARY KEY, nombre TEXT, fecha_nacimiento TEXT, seccion TEXT,
            calle TEXT, num_ext TEXT, num_int TEXT, cp TEXT, fraccionamiento TEXT, telefono TEXT,
            folio_tarjeta TEXT, foto_frente TEXT, foto_reverso TEXT
        )
    ''')
    
    # 2. Tabla de Usuarios (Agregamos session_token para control de sesiones)
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS usuarios (
            id SERIAL PRIMARY KEY, username TEXT UNIQUE, password TEXT, rol TEXT, session_token TEXT
        )
    ''')
    # Intenta agregar la columna si la tabla ya existía antes de esta actualización
    try:
        cursor.execute("ALTER TABLE usuarios ADD COLUMN session_token TEXT")
    except psycopg2.errors.DuplicateColumn:
        conn.rollback()
    
    # 3. Tabla de Auditoría (Movimientos)
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS auditoria (
            id SERIAL PRIMARY KEY, usuario TEXT, accion TEXT, detalles TEXT, 
            fecha TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    
    # 4. Crear usuarios por defecto si no hay ninguno
    cursor.execute("SELECT COUNT(*) FROM usuarios")
    if cursor.fetchone()[0] == 0:
        cursor.execute("INSERT INTO usuarios (username, password, rol) VALUES (%s, %s, %s)", ("admin", hash_password("admin123"), "admin"))
        cursor.execute("INSERT INTO usuarios (username, password, rol) VALUES (%s, %s, %s)", ("soporte", hash_password("soporte123"), "soporte"))
        cursor.execute("INSERT INTO usuarios (username, password, rol) VALUES (%s, %s, %s)", ("captura", hash_password("captura123"), "capturador"))
                       
    conn.commit()
    conn.close()

@app.on_event("startup")
def startup():
    if DATABASE_URL:
        init_db()

# --- FUNCIÓN DE AUDITORÍA ---
def registrar_movimiento(usuario: str, accion: str, detalles: str):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("INSERT INTO auditoria (usuario, accion, detalles) VALUES (%s, %s, %s)", (usuario, accion, detalles))
    conn.commit()
    conn.close()

# --- SISTEMA DE LOGIN Y SESIONES ESTRICTAS ---
def get_current_user(request: Request):
    user_session = request.session.get("user")
    if not user_session: return None
    
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=RealDictCursor)
    cursor.execute("SELECT rol, session_token FROM usuarios WHERE username = %s", (user_session["username"],))
    db_user = cursor.fetchone()
    conn.close()
    
    # Si el token no coincide (ej. el admin cerró su sesión), se invalida
    if not db_user or db_user["session_token"] != user_session.get("token"):
        request.session.clear()
        return None
    
    user_session["rol"] = db_user["rol"]
    return user_session

@app.get("/login", response_class=HTMLResponse)
async def login_get(request: Request):
    return templates.TemplateResponse(request=request, name="login.html")

@app.post("/login")
async def login_post(request: Request, username: str = Form(...), password: str = Form(...)):
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=RealDictCursor)
    cursor.execute("SELECT * FROM usuarios WHERE username = %s AND password = %s", (username, hash_password(password)))
    user = cursor.fetchone()
    
    if user:
        # Generar nuevo token de sesión
        new_token = str(uuid.uuid4())
        cursor.execute("UPDATE usuarios SET session_token = %s WHERE id = %s", (new_token, user["id"]))
        conn.commit()
        conn.close()
        
        request.session["user"] = {"username": user["username"], "rol": user["rol"], "token": new_token}
        registrar_movimiento(user["username"], "Inicio de Sesión", "El usuario ingresó al sistema.")
        return RedirectResponse(url="/", status_code=303)
    
    conn.close()
    return templates.TemplateResponse(request=request, name="login.html", context={"error": "Credenciales incorrectas"})

@app.get("/logout")
async def logout(request: Request):
    user = get_current_user(request)
    if user: registrar_movimiento(user["username"], "Cierre de Sesión", "El usuario salió del sistema.")
    request.session.clear()
    return RedirectResponse(url="/login", status_code=303)

# --- RUTAS PRINCIPALES ---
@app.get("/", response_class=HTMLResponse)
async def inicio(request: Request):
    user = get_current_user(request)
    if not user: return RedirectResponse(url="/login", status_code=303)
    return templates.TemplateResponse(request=request, name="index.html", context={"user": user})

@app.get("/registros", response_class=HTMLResponse)
async def ver_registros(request: Request):
    user = get_current_user(request)
    if not user or user["rol"] not in ["admin", "soporte"]:
        return RedirectResponse(url="/", status_code=303)
        
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=RealDictCursor)
    cursor.execute("SELECT * FROM beneficiarios ORDER BY id DESC")
    registros = cursor.fetchall()
    conn.close()
    return templates.TemplateResponse(request=request, name="registros.html", context={"registros": registros, "user": user})

@app.post("/agregar")
async def agregar_registro(
    request: Request, nombre: str = Form(...), fecha_nacimiento: str = Form(...), seccion: str = Form(...), calle: str = Form(...), 
    num_ext: str = Form(...), num_int: str = Form(""), cp: str = Form(...), fraccionamiento: str = Form(...),
    telefono: str = Form(...), folio_tarjeta: str = Form(...), foto_frente: str = Form(""), foto_reverso: str = Form("")
):
    user = get_current_user(request)
    if not user: return JSONResponse(content={"status": "error", "message": "Sesión expirada"}, status_code=401)

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('''
        INSERT INTO beneficiarios (nombre, fecha_nacimiento, seccion, calle, num_ext, num_int, cp, fraccionamiento, telefono, folio_tarjeta, foto_frente, foto_reverso) 
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s) RETURNING id
    ''', (nombre, fecha_nacimiento, seccion, calle, num_ext, num_int, cp, fraccionamiento, telefono, folio_tarjeta, foto_frente, foto_reverso))
    new_id = cursor.fetchone()[0]
    conn.commit()
    conn.close()
    
    registrar_movimiento(user["username"], "Captura", f"Registró al beneficiario: {nombre} (ID: {new_id})")
    return JSONResponse(content={"status": "success", "message": "Registro guardado correctamente"})

@app.post("/editar/{reg_id}")
async def editar_registro(
    request: Request, reg_id: int, nombre: str = Form(...), fecha_nacimiento: str = Form(...), 
    seccion: str = Form(...), calle: str = Form(...), num_ext: str = Form(...), num_int: str = Form(""),
    cp: str = Form(...), fraccionamiento: str = Form(...), telefono: str = Form(...), folio_tarjeta: str = Form(...)
):
    user = get_current_user(request)
    if not user or user["rol"] not in ["admin", "soporte"]:
        return JSONResponse(content={"status": "error", "message": "Acceso denegado"}, status_code=403)
        
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('''
        UPDATE beneficiarios SET nombre=%s, fecha_nacimiento=%s, seccion=%s, calle=%s, num_ext=%s, num_int=%s, cp=%s, fraccionamiento=%s, telefono=%s, folio_tarjeta=%s 
        WHERE id=%s
    ''', (nombre, fecha_nacimiento, seccion, calle, num_ext, num_int, cp, fraccionamiento, telefono, folio_tarjeta, reg_id))
    conn.commit()
    conn.close()
    
    registrar_movimiento(user["username"], "Edición", f"Editó datos del beneficiario: {nombre} (ID: {reg_id})")
    return JSONResponse(content={"status": "success"})

@app.post("/eliminar/{reg_id}")
async def eliminar_registro(request: Request, reg_id: int):
    user = get_current_user(request)
    if not user or user["rol"] not in ["admin", "soporte"]:
        return JSONResponse(content={"status": "error", "message": "Acceso denegado"}, status_code=403)
        
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=RealDictCursor)
    cursor.execute("SELECT nombre FROM beneficiarios WHERE id = %s", (reg_id,))
    beneficiario = cursor.fetchone()
    
    if beneficiario:
        cursor.execute("DELETE FROM beneficiarios WHERE id = %s", (reg_id,))
        conn.commit()
        registrar_movimiento(user["username"], "Eliminación", f"Eliminó al beneficiario: {beneficiario['nombre']} (ID: {reg_id})")
        
    conn.close()
    return JSONResponse(content={"status": "success"})

@app.get("/descargar_excel")
async def descargar_excel(request: Request):
    user = get_current_user(request)
    if not user or user["rol"] not in ["admin", "soporte"]:
        return RedirectResponse(url="/", status_code=303)
        
    registrar_movimiento(user["username"], "Exportación", "Descargó la base de datos en Excel")
    
    conn = get_db_connection()
    df = pd.read_sql_query('SELECT id AS "ID", folio_tarjeta AS "Folio", nombre AS "Nombre", fecha_nacimiento AS "Nacimiento", seccion AS "Sección", calle AS "Calle", num_ext AS "Num_Ext", num_int AS "Num_Int", fraccionamiento AS "Fraccionamiento", cp AS "CP", telefono AS "Teléfono" FROM beneficiarios', conn)
    conn.close()
    
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Beneficiarios')
        worksheet = writer.sheets['Beneficiarios']
        guinda_fill = PatternFill(start_color="691C32", end_color="691C32", fill_type="solid")
        white_font = Font(color="FFFFFF", bold=True, size=11)
        thin_border = Border(left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD'), top=Side(style='thin', color='DDDDDD'), bottom=Side(style='thin', color='DDDDDD'))
        worksheet.freeze_panes = 'A2'
        for col_num, cell in enumerate(worksheet[1], 1):
            cell.fill = guinda_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        for idx, col in enumerate(worksheet.columns, 1):
            max_length = max((len(str(cell.value)) for cell in col if cell.value), default=0)
            worksheet.column_dimensions[get_column_letter(idx)].width = min(max_length + 3, 40)
            
    buffer.seek(0)
    return StreamingResponse(buffer, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="Base_Datos_Bienestar.xlsx"'})

# --- GESTIÓN DE USUARIOS (SOLO ADMIN) ---
@app.get("/usuarios", response_class=HTMLResponse)
async def ver_usuarios(request: Request):
    user = get_current_user(request)
    if not user or user["rol"] != "admin": return RedirectResponse(url="/", status_code=303)
        
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=RealDictCursor)
    cursor.execute("SELECT id, username, rol, CASE WHEN session_token IS NOT NULL THEN 'En línea' ELSE 'Desconectado' END as estado FROM usuarios ORDER BY id ASC")
    usuarios_db = cursor.fetchall()
    conn.close()
    return templates.TemplateResponse(request=request, name="usuarios.html", context={"user": user, "usuarios_db": usuarios_db})

@app.post("/usuarios/crear")
async def crear_usuario(request: Request, username: str = Form(...), password: str = Form(...), rol: str = Form(...)):
    user = get_current_user(request)
    if not user or user["rol"] != "admin": return JSONResponse(content={"status": "error"}, status_code=403)
    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("INSERT INTO usuarios (username, password, rol) VALUES (%s, %s, %s)", (username, hash_password(password), rol))
        conn.commit()
        registrar_movimiento(user["username"], "Gestión Personal", f"Creó al usuario '{username}' con rol '{rol}'")
        return JSONResponse(content={"status": "success", "message": "Usuario creado exitosamente"})
    except psycopg2.IntegrityError:
        conn.rollback()
        return JSONResponse(content={"status": "error", "message": "El usuario ya existe"})
    finally:
        conn.close()

@app.post("/usuarios/password/{id}")
async def cambiar_password(request: Request, id: int, new_password: str = Form(...)):
    user = get_current_user(request)
    if not user or user["rol"] != "admin": return JSONResponse(content={"status": "error"}, status_code=403)
    conn = get_db_connection()
    cursor = conn.cursor()
    # Cambia clave y BORRA el token para forzar cierre de sesión
    cursor.execute("UPDATE usuarios SET password = %s, session_token = NULL WHERE id = %s RETURNING username", (hash_password(new_password), id))
    target = cursor.fetchone()[0]
    conn.commit()
    conn.close()
    registrar_movimiento(user["username"], "Seguridad", f"Cambió la contraseña y cerró sesión de '{target}'")
    return JSONResponse(content={"status": "success"})

@app.post("/usuarios/eliminar/{id}")
async def eliminar_usuario(request: Request, id: int):
    user = get_current_user(request)
    if not user or user["rol"] != "admin": return JSONResponse(content={"status": "error"}, status_code=403)
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM usuarios WHERE id = %s RETURNING username", (id,))
    target = cursor.fetchone()[0]
    conn.commit()
    conn.close()
    registrar_movimiento(user["username"], "Gestión Personal", f"Eliminó al usuario '{target}'")
    return JSONResponse(content={"status": "success"})

@app.post("/usuarios/logout/{id}")
async def forzar_logout(request: Request, id: int):
    user = get_current_user(request)
    if not user or user["rol"] != "admin": return JSONResponse(content={"status": "error"}, status_code=403)
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("UPDATE usuarios SET session_token = NULL WHERE id = %s RETURNING username", (id,))
    target = cursor.fetchone()[0]
    conn.commit()
    conn.close()
    registrar_movimiento(user["username"], "Seguridad", f"Forzó el cierre de sesión de '{target}'")
    return JSONResponse(content={"status": "success"})

# --- AUDITORÍA (SOLO ADMIN) ---
@app.get("/auditoria", response_class=HTMLResponse)
async def ver_auditoria(request: Request):
    user = get_current_user(request)
    if not user or user["rol"] != "admin": return RedirectResponse(url="/", status_code=303)
        
    conn = get_db_connection()
    cursor = conn.cursor(cursor_factory=RealDictCursor)
    cursor.execute("SELECT * FROM auditoria ORDER BY id DESC LIMIT 200") # Muestra últimos 200 para no saturar
    logs = cursor.fetchall()
    conn.close()
    return templates.TemplateResponse(request=request, name="auditoria.html", context={"user": user, "logs": logs})